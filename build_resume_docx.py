"""Build Mrunal_Kubal_Java_TechLead_GenAI.docx from scratch using python-docx."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── Colour palette ──────────────────────────────────────────────────────────
DARK_BLUE = RGBColor(0x1a, 0x3c, 0x6e)
MID_BLUE  = RGBColor(0x2c, 0x5f, 0x9e)
BLACK     = RGBColor(0x1a, 0x1a, 0x1a)
GREY      = RGBColor(0x44, 0x44, 0x44)

doc = Document()

# ── Page margins ─────────────────────────────────────────────────────────────
for section in doc.sections:
    section.top_margin    = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin   = Cm(2.0)
    section.right_margin  = Cm(2.0)

# ── Helper: set paragraph spacing ────────────────────────────────────────────
def para_space(para, before=0, after=0, line=None):
    pf = para.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after  = Pt(after)
    if line:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(line)

# ── Helper: add a run with font settings ────────────────────────────────────
def add_run(para, text, bold=False, italic=False, size=10, color=BLACK, underline=False):
    run = para.add_run(text)
    run.bold      = bold
    run.italic    = italic
    run.underline = underline
    run.font.size  = Pt(size)
    run.font.color.rgb = color
    run.font.name  = "Calibri"
    return run

# ── Helper: add a bottom-border (HR) to a paragraph ─────────────────────────
def add_bottom_border(para, color="1A3C6E", size="6"):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"),   "single")
    bottom.set(qn("w:sz"),    size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)

# ── Helper: section title ─────────────────────────────────────────────────────
def section_title(text):
    p = doc.add_paragraph()
    para_space(p, before=8, after=2)
    add_bottom_border(p)
    add_run(p, text.upper(), bold=True, size=10.5, color=DARK_BLUE)
    return p

# ── Helper: bullet ────────────────────────────────────────────────────────────
def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style="List Bullet")
    para_space(p, before=1, after=1, line=12.5)
    p.paragraph_format.left_indent  = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    if bold_prefix:
        add_run(p, bold_prefix, bold=True, size=9.5, color=BLACK)
        remaining = text[len(bold_prefix):]
        add_run(p, remaining, size=9.5, color=BLACK)
    else:
        add_run(p, text, size=9.5, color=BLACK)
    return p

# ═══════════════════════════════════════════════════════════════════════════════
#  HEADER
# ═══════════════════════════════════════════════════════════════════════════════
name_p = doc.add_paragraph()
name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_space(name_p, before=0, after=2)
add_run(name_p, "MRUNAL KUBAL", bold=True, size=22, color=DARK_BLUE)
add_bottom_border(name_p, size="12")

tagline_p = doc.add_paragraph()
tagline_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_space(tagline_p, before=3, after=2)
add_run(tagline_p,
        "Java Tech Lead  |  GenAI Developer Productivity  |  Enterprise Architect",
        italic=True, size=10, color=MID_BLUE)

contact_p = doc.add_paragraph()
contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para_space(contact_p, before=2, after=6)
add_run(contact_p,
        "Thane, Maharashtra  |  +91 8879756218 / +91 8108246077  |  mrunalkubal20@gmail.com",
        size=9.5, color=GREY)

# ═══════════════════════════════════════════════════════════════════════════════
#  PROFESSIONAL SUMMARY
# ═══════════════════════════════════════════════════════════════════════════════
section_title("Professional Summary")
summary_p = doc.add_paragraph()
para_space(summary_p, before=2, after=3, line=13)
summary_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
add_run(summary_p,
        "Results-driven Java Tech Lead with 9+ years of enterprise software engineering experience, "
        "specialising in GenAI-powered developer productivity tools, enterprise Java/J2EE architecture, "
        "and cloud-native solutions. Proven record leading cross-functional teams, defining scalable "
        "microservices architectures, and integrating Large Language Models (LLMs) — including "
        "OpenAI/ChatGPT, LangChain, and Perplexity — into production enterprise workflows. Deep expertise "
        "in Spring Boot, Kubernetes, Docker, Azure, and CI/CD pipeline automation. Experienced in building "
        "AI-assisted migration and code-generation tooling, with strong working knowledge of enterprise "
        "security, GDPR compliance, and governance standards across Banking, Insurance, and FinTech domains.",
        size=10, color=BLACK)

# ═══════════════════════════════════════════════════════════════════════════════
#  CORE TECHNICAL SKILLS  (2-col table)
# ═══════════════════════════════════════════════════════════════════════════════
section_title("Core Technical Skills")

skills = [
    ("Java Architecture",
     "Java 8/11/17, J2EE, Spring Boot, Spring MVC, Spring Security, JPA, Hibernate, REST APIs, "
     "Microservices, API Gateway, Design Patterns, Enterprise Architecture"),
    ("GenAI & LLMs",
     "OpenAI/ChatGPT, LangChain, HuggingFace, Perplexity API, Prompt Engineering, Agentic AI, "
     "RAG pipelines, LLM-powered code generation & migration tools"),
    ("IDE & Tooling",
     "IntelliJ IDEA, Eclipse, IntelliJ Platform SDK / Eclipse PDE, Developer productivity tooling, "
     "VS Code extensions"),
    ("Cloud & DevOps",
     "Azure (Key Vault, AKS, App Services), AWS (EC2, S3, Lambda), GCP fundamentals, Docker, "
     "Kubernetes, Jenkins, GitHub Actions, CI/CD pipeline design"),
    ("Security & Compliance",
     "GDPR / Data Privacy, RBAC, OAuth 2.0, Enterprise security standards, Audit logging, "
     "Azure Key Vault, Secrets Management"),
    ("Frontend",
     "Angular 8+, React, TypeScript, JavaScript, HTML5, CSS3"),
    ("Databases",
     "Oracle, MySQL, MongoDB"),
    ("Testing & Quality",
     "JUnit, Mockito, Jest, TDD, Code Reviews, SonarQube"),
    ("Observability",
     "Grafana, Prometheus, Swagger / OpenAPI"),
    ("Methodology",
     "Agile / Scrum, Maven, Git, Architecture & Sequence Diagramming, Enterprise Design Reviews"),
]

tbl = doc.add_table(rows=len(skills), cols=2)
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl.style = "Table Grid"

# Remove all borders from the table
def remove_table_borders(table):
    tbl_el = table._tbl
    tblPr = tbl_el.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl_el.insert(0, tblPr)
    tblBorders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        b.set(qn("w:val"), "none")
        tblBorders.append(b)
    tblPr.append(tblBorders)

remove_table_borders(tbl)

# Column widths
tbl.columns[0].width = Inches(1.55)
tbl.columns[1].width = Inches(5.45)

for i, (label, value) in enumerate(skills):
    row = tbl.rows[i]
    row.height = None

    c0 = row.cells[0]
    c1 = row.cells[1]

    c0.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    c1.vertical_alignment = WD_ALIGN_VERTICAL.TOP

    p0 = c0.paragraphs[0]
    para_space(p0, before=1, after=1, line=12)
    add_run(p0, label, bold=True, size=9.5, color=DARK_BLUE)

    p1 = c1.paragraphs[0]
    para_space(p1, before=1, after=1, line=12)
    add_run(p1, value, size=9.5, color=BLACK)

# ═══════════════════════════════════════════════════════════════════════════════
#  PROFESSIONAL EXPERIENCE
# ═══════════════════════════════════════════════════════════════════════════════
section_title("Professional Experience")

jobs = [
    {
        "company": "Virtusa",
        "dates": "Feb 2022 – Present",
        "title": "Lead Software Developer → Java Tech Lead",
        "tech": "Java 17, Spring Boot, Azure (AKS/Key Vault), Microservices, Docker, Kubernetes, CI/CD (GitHub Actions/Jenkins), GenAI/LLM Integration",
        "bullets": [
            ("Led enterprise architecture", " for a large-scale Credit Card Platform, defining microservices blueprints, API contracts, and governance standards adopted across the engineering org."),
            ("Architected and delivered a GenAI-powered Angular-to-React migration assistant", " using OpenAI ChatGPT and Perplexity API, reducing manual migration effort by ~40% for enterprise frontend codebases."),
            ("Integrated LLM-driven code-generation tooling", " into the developer workflow (IntelliJ-compatible prompts, CI/CD hooks), enabling automated boilerplate generation and documentation drafts."),
            ("Spearheaded migration of legacy Backbase services to Azure Cloud", ", improving horizontal scalability; integrated Azure Key Vault for enterprise-grade secrets management."),
            ("Designed and enforced automated CI/CD pipelines", " (GitHub Actions + Jenkins) including security-gate stages, SonarQube analysis, and container vulnerability scanning aligned with enterprise compliance policies."),
            ("Improved critical API response times by 30%", " through profiling, query optimisation, and caching strategies."),
            ("Produced architecture, sequence, and deployment diagrams", " for CTO-level stakeholder reviews; acted as primary technical authority for design approvals."),
            ("Mentored junior and mid-level engineers", " through structured code reviews, architecture walkthroughs, and pair-programming sessions; established team coding standards and best-practice playbooks."),
            ("Collaborated with product, security, and platform engineering teams", " to align GenAI tooling with corporate data-privacy (GDPR) and auditability requirements."),
        ],
    },
    {
        "company": "Capgemini",
        "dates": "Mar 2020 – Jan 2022",
        "title": "Senior Software Developer",
        "tech": "Java 11, Spring Boot, Angular, Microservices, REST APIs, Docker",
        "bullets": [
            ("Developed and owned dynamic insurance product modules", " (Travel, Car, Cross-sell) via an API-first microservices approach integrated with third-party enterprise systems."),
            ("Designed modular, reusable service components", " to accelerate developer productivity and reduce time-to-feature across teams."),
            ("Optimised Angular frontend load time by 25%", " through lazy loading, change-detection strategy tuning, and bundle analysis."),
            ("Achieved 85%+ Jest unit test coverage", " for critical business components; promoted TDD practices within the team."),
            ("Contributed to API governance reviews", " ensuring OpenAPI specifications met enterprise integration standards."),
        ],
    },
    {
        "company": "CGI",
        "dates": "Dec 2018 – Feb 2020",
        "title": "Senior Software Developer",
        "tech": "Java 8, Spring MVC, MySQL, REST APIs, Enterprise Credit Systems",
        "bullets": [
            ("Built core credit-decisioning modules", " for the STRATA Credit Decision Engine, deployed at major U.S. banking institutions."),
            ("Delivered a 20% performance improvement", " through client/server logic optimisation and strategic query redesign."),
            ("Contributed to architectural design", " of large-scale credit workflows, ensuring extensibility and compliance with financial industry standards."),
        ],
    },
    {
        "company": "C-EDGE Technologies",
        "dates": "Feb 2017 – Dec 2018",
        "title": "Assistant System Analyst",
        "tech": "Java 8, Spring MVC, Oracle, Multithreading, ISO 8583, SFMS",
        "bullets": [
            ("Implemented RBAC", " (Maker/Checker/Auditor/Admin) for IBK User Interface Portal, aligning with enterprise security governance and audit requirements."),
            ("Integrated NEFT/RTGS processing with SFMS", " including acknowledgements, reporting, and UTR enquiry modules — ensuring compliance with RBI regulatory standards."),
            ("Implemented ISO 8583 message processing", " for KIOSK, Micro ATM, and M-Passbook interfaces — real-time financial transaction handling at scale."),
            ("Built bulk upload and audit reporting modules", " supporting regulatory compliance and operational traceability."),
        ],
    },
    {
        "company": "Unotech Software Pvt. Ltd.",
        "dates": "Jun 2016 – Feb 2017",
        "title": "Software Engineer",
        "tech": "Node.js, MongoDB, AWS (EC2, S3)",
        "bullets": [
            ("Developed RESTful web services", " on Node.js + MongoDB for mobile applications hosted on AWS, gaining early exposure to cloud-native deployments and push-notification pipelines."),
        ],
    },
]

for job in jobs:
    # Company + dates on one line
    p_co = doc.add_paragraph()
    para_space(p_co, before=6, after=0)
    add_run(p_co, job["company"], bold=True, size=10.5, color=DARK_BLUE)
    tab_spaces = "\t"
    r_date = p_co.add_run(tab_spaces + job["dates"])
    r_date.italic = True
    r_date.font.size = Pt(9)
    r_date.font.color.rgb = GREY
    r_date.font.name = "Calibri"
    p_co.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)

    # Title
    p_title = doc.add_paragraph()
    para_space(p_title, before=1, after=0)
    add_run(p_title, job["title"], italic=True, size=10, color=BLACK)

    # Tech stack
    p_tech = doc.add_paragraph()
    para_space(p_tech, before=1, after=2)
    add_run(p_tech, job["tech"], size=9, color=GREY, italic=True)

    # Bullets
    for bold_part, rest in job["bullets"]:
        p = doc.add_paragraph(style="List Bullet")
        para_space(p, before=1, after=1, line=12.5)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        add_run(p, bold_part, bold=True, size=9.5, color=BLACK)
        add_run(p, rest, size=9.5, color=BLACK)

# ═══════════════════════════════════════════════════════════════════════════════
#  KEY GENAI & DEVELOPER PRODUCTIVITY PROJECTS
# ═══════════════════════════════════════════════════════════════════════════════
section_title("Key GenAI & Developer Productivity Projects")

projects = [
    {
        "name": "GenAI Angular-to-React Migration Assistant",
        "dates": "2023 – Present",
        "tech": "OpenAI ChatGPT, Perplexity API, LangChain, Java/Spring Boot, IntelliJ plugin integration",
        "bullets": [
            ("Designed an LLM-powered assistant", " that analyses Angular component trees, infers business logic, and generates idiomatic React equivalents — integrated as an IntelliJ-compatible developer tool."),
            ("Implemented prompt engineering pipelines", " with contextual chunking and RAG-style retrieval to handle large enterprise codebases exceeding 100K LOC."),
            ("Reduced average migration time per component by ~40%", ", providing measurable ROI for enterprise-wide frontend modernisation initiatives."),
        ],
    },
    {
        "name": "AI-Augmented CI/CD Pipeline — Code Quality & Governance Gates",
        "dates": "2024",
        "tech": "LangChain, OpenAI, GitHub Actions, Docker, Azure AKS",
        "bullets": [
            ("Designed and prototyped an LLM-powered code-review gate", " in CI/CD pipelines that flags security vulnerabilities, deviations from architecture guidelines, and GDPR data-handling anti-patterns before merge."),
            ("Integrated results as inline PR annotations", " within the enterprise developer workflow, driving adoption without disrupting existing toolchain habits."),
        ],
    },
]

for proj in projects:
    p_co = doc.add_paragraph()
    para_space(p_co, before=6, after=0)
    add_run(p_co, proj["name"], bold=True, size=10.5, color=DARK_BLUE)
    r_date = p_co.add_run("\t" + proj["dates"])
    r_date.italic = True
    r_date.font.size = Pt(9)
    r_date.font.color.rgb = GREY
    r_date.font.name = "Calibri"
    p_co.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_ALIGN_PARAGRAPH.RIGHT)

    p_tech = doc.add_paragraph()
    para_space(p_tech, before=1, after=2)
    add_run(p_tech, proj["tech"], size=9, color=GREY, italic=True)

    for bold_part, rest in proj["bullets"]:
        p = doc.add_paragraph(style="List Bullet")
        para_space(p, before=1, after=1, line=12.5)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.15)
        add_run(p, bold_part, bold=True, size=9.5, color=BLACK)
        add_run(p, rest, size=9.5, color=BLACK)

# ═══════════════════════════════════════════════════════════════════════════════
#  EDUCATION
# ═══════════════════════════════════════════════════════════════════════════════
section_title("Education")

edu = [
    ("B.E. in Computer Engineering – Shah & Anchor Kutchhi Engineering College", "2015"),
    ("Diploma in Computer Engineering – Vidyalankar Polytechnic", "2012"),
]
for degree, year in edu:
    p = doc.add_paragraph()
    para_space(p, before=2, after=1)
    add_run(p, degree, bold=True, size=9.5, color=BLACK)
    add_run(p, "   " + year, italic=True, size=9, color=GREY)

# ═══════════════════════════════════════════════════════════════════════════════
#  AWARDS & PROFESSIONAL DEVELOPMENT  (2-col table)
# ═══════════════════════════════════════════════════════════════════════════════
section_title("Awards & Professional Development")

aw_tbl = doc.add_table(rows=1, cols=2)
aw_tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
remove_table_borders(aw_tbl)
aw_tbl.columns[0].width = Inches(3.5)
aw_tbl.columns[1].width = Inches(3.5)

c_left  = aw_tbl.rows[0].cells[0]
c_right = aw_tbl.rows[0].cells[1]

def awards_block(cell, heading, items):
    p = cell.paragraphs[0]
    para_space(p, before=2, after=2)
    add_run(p, heading, bold=True, size=9.5, color=BLACK)
    for item in items:
        p2 = cell.add_paragraph(style="List Bullet")
        para_space(p2, before=1, after=1, line=12)
        p2.paragraph_format.left_indent = Inches(0.2)
        p2.paragraph_format.first_line_indent = Inches(-0.15)
        add_run(p2, item, size=9.5, color=BLACK)

awards_block(c_left,  "Awards & Recognition",       ["Best Performer of the Year", "Pat on the Back Award"])
awards_block(c_right, "Certifications & Trainings", ["Generative AI — In Progress", "Agentic AI", "Docker", "Kubernetes"])

# ── Save ──────────────────────────────────────────────────────────────────────
out = "/workspace/Mrunal_Kubal_Java_TechLead_GenAI.docx"
doc.save(out)
print(f"Saved: {out}")
